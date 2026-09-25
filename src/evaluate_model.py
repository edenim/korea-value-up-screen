"""Evaluate every formula in the Excel model and export labeled results.

openpyxl writes formulas without cached values, so this script computes them with the
`formulas` package, fails loudly on any formula error, and writes
outputs/model_values.json: {sheet: {row label: {column header: value}}}.
The initiation note reads its numbers from that file, so every figure traces to the model.

Run from the repo root:  .venv/bin/python src/evaluate_model.py
"""

import json

import formulas
import numpy as np
import openpyxl

from config import OUTPUT_DIR, ROOT

MODEL_PATH = ROOT / "model" / "Navien_009450_model.xlsx"
OUT_PATH = OUTPUT_DIR / "model_values.json"
HEADER_ROW = 4


def evaluate():
    xl = formulas.ExcelModel().loads(str(MODEL_PATH)).finish()
    solution = xl.calculate()
    values, errors = {}, []
    for ref, cell in solution.items():
        ref = str(ref)
        if "!" not in ref or ":" in ref.split("!")[1]:
            continue
        sheet, addr = ref.split("]")[1].split("!")
        value = cell.value[0, 0] if hasattr(cell, "value") else cell
        if isinstance(value, formulas.tokens.operand.XlError):
            errors.append(ref)
            continue
        if isinstance(value, (np.floating, np.integer)):
            value = float(value)
        values[(sheet.strip("'").upper(), addr.upper())] = value
    if errors:
        raise RuntimeError(f"{len(errors)} formula errors, e.g. {errors[:5]}")
    return values


def labeled(values):
    """Map values to {sheet: {row label: {column header: value}}} using column A and row 4."""
    wb = openpyxl.load_workbook(MODEL_PATH)
    out = {}
    for ws in wb.worksheets:
        sheet = {}
        headers = {c.column_letter: c.value for c in ws[HEADER_ROW] if c.value}
        for row in ws.iter_rows(min_row=HEADER_ROW + 1):
            lab = row[0].value
            # Rows without a text label (sensitivity grids) are keyed by their row number
            row_key = lab.strip() if isinstance(lab, str) else f"row{row[0].row}"
            entry = {}
            if not isinstance(lab, str) and lab is not None:
                entry["A"] = values.get((ws.title.upper(), row[0].coordinate), lab)
            for cell in row[1:]:
                key = (ws.title.upper(), cell.coordinate)
                val = values.get(key, cell.value)
                if val is None or isinstance(val, str) and val.startswith("="):
                    continue
                col_name = headers.get(cell.column_letter, cell.column_letter)
                entry[col_name] = val
            if entry or isinstance(lab, str):
                sheet[row_key] = entry
        out[ws.title] = sheet
    return out


def main():
    values = evaluate()
    data = labeled(values)
    OUT_PATH.write_text(json.dumps(data, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    dcf = data["DCF"]
    print(f"Evaluated {len(values)} cells, 0 errors.")
    print("Value per share:", round(dcf["Equity value per share (KRW)"]["Value"]),
          "| upside:", round(dcf["Upside / (downside)"]["Value"], 4), "| rating:", dcf["Rating"]["Value"])


if __name__ == "__main__":
    main()
